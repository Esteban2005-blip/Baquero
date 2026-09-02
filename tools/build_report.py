from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Informe_tecnico_Baquero_Notes.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "182033"
MUTED = "59657A"
PRIMARY = "4338CA"
PRIMARY_SOFT = "E8E7FF"
LIGHT_GRAY = "F2F4F7"
OUTLINE = "C8CFDC"
SUCCESS = "087A55"
WARNING = "9A4D00"
ERROR = "B42318"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    total = sum(widths_dxa)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:cantSplit")
    tr_pr.append(marker)


def set_font(run, name="Calibri", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        run._r.append(element)
    set_font(run, size=9, color=MUTED)


def add_caption(doc, text):
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_font(run, size=9, italic=True, color=MUTED)
    return paragraph


def add_callout(doc, label, text, accent=PRIMARY, fill=PRIMARY_SOFT):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(8)
    paragraph.paragraph_format.right_indent = Pt(8)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:space"), "4")
        border.set(qn("w:color"), accent)
        borders.append(border)
    p_pr.append(borders)
    label_run = paragraph.add_run(f"{label}: ")
    set_font(label_run, bold=True, color=accent)
    text_run = paragraph.add_run(text)
    set_font(text_run, color=INK)


def add_table(doc, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    repeat_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_GRAY)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(str(header))
        set_font(run, size=font_size, bold=True, color=INK)
    prevent_row_split(table.rows[0])
    for row_data in rows:
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        for index, value in enumerate(row_data):
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            set_font(run, size=font_size, color=INK)
    set_table_geometry(table, widths)
    return table


def add_labeled_paragraph(doc, label, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    label_run = paragraph.add_run(f"{label}: ")
    set_font(label_run, bold=True, color=DARK_BLUE)
    value_run = paragraph.add_run(text)
    set_font(value_run, color=INK)
    return paragraph


def add_code_block(doc, path):
    lines = path.read_text(encoding="utf-8").splitlines()
    for number, line in enumerate(lines, start=1):
        paragraph = doc.add_paragraph(style="Code")
        paragraph.paragraph_format.keep_together = False
        shade_paragraph(paragraph, "F7F8FA")
        run = paragraph.add_run(f"{number:03d} | {line}")
        set_font(run, name="Consolas", size=7.2, color=INK)


def set_image_alt(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("name", title)
    doc_pr.set("descr", description)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    code = styles.add_style("Code", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(7.2)
    code.font.color.rgb = RGBColor.from_string(INK)
    code.paragraph_format.left_indent = Inches(0.08)
    code.paragraph_format.right_indent = Inches(0.08)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    code.paragraph_format.line_spacing = Pt(8.6)

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)


def configure_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True
    doc.settings.odd_and_even_pages_header_footer = True

    for footer in (section.footer, section.even_page_footer):
        paragraph = footer.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run("Baquero Notes  |  Pagina ")
        set_font(run, size=8.5, color=MUTED)
        add_page_field(paragraph)


def add_cover(doc):
    for _ in range(4):
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(12)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    run = kicker.add_run("INFORME TECNICO")
    set_font(run, size=11, bold=True, color=PRIMARY)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("Sistema de componentes reutilizables\npara Baquero Notes")
    set_font(run, size=28, bold=True, color=INK)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(52)
    run = subtitle.add_run("Flutter + API Flask | Diseño, implementación y verificación")
    set_font(run, size=14, color=MUTED)

    metadata = [
        "Autor: [COMPLETAR NOMBRE]",
        "Asignatura / curso: [COMPLETAR]",
        "Docente: [COMPLETAR]",
        "Institución: [COMPLETAR]",
        "Fecha: 23 de agosto de 2026",
    ]
    for value in metadata:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(value)
        set_font(run, size=10.5, color=INK, bold=value.startswith("Autor"))

    doc.add_page_break()


def build():
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    add_cover(doc)

    doc.add_heading("Resumen ejecutivo", level=1)
    doc.add_paragraph(
        "Se reorganizó Baquero Notes como una aplicación Flutter conectada a una API Flask segura. "
        "La solución separa tokens, componentes, pantallas y acceso HTTP; incorpora un catálogo vivo de cuatro "
        "componentes reutilizables y resuelve de forma explícita carga, vacío y error en la pantalla de notas."
    )
    add_callout(
        doc,
        "Resultado",
        "El flujo de backend registro - crear - listar - actualizar - eliminar fue verificado sobre una base SQLite aislada. "
        "La interfaz quedó preparada para pruebas Flutter, pero el entorno de elaboración no incluía Flutter SDK; esta limitación se declara y no se presenta como una prueba ejecutada.",
        accent=SUCCESS,
        fill="E8F5EF",
    )

    doc.add_heading("1. Inventario de API y pantallas", level=1)
    doc.add_paragraph(
        "La API es la única fuente de datos del flujo nuevo. Los componentes visuales no conocen rutas, tokens JWT ni reglas de persistencia; "
        "las pantallas coordinan acciones y delegan las solicitudes a ApiService."
    )
    doc.add_heading("1.1 Endpoints disponibles", level=2)
    endpoint_rows = [
        ("POST", "/api/auth/register", "Registrar usuario y emitir tokens", "RegisterPage", "Consumido"),
        ("POST", "/api/auth/login", "Autenticar y emitir tokens", "LoginPage", "Consumido"),
        ("POST", "/api/auth/refresh", "Renovar token de acceso", "Sin pantalla", "Disponible"),
        ("POST", "/api/auth/logout", "Revocar token", "NotesPage", "Consumido"),
        ("GET", "/api/notes", "Listar notas paginadas", "NotesPage", "Consumido"),
        ("POST", "/api/notes", "Crear nota", "NoteEditorDialog", "Consumido"),
        ("GET", "/api/notes/{id}", "Consultar una nota", "Sin pantalla dedicada", "Disponible"),
        ("PUT", "/api/notes/{id}", "Actualizar una nota", "NoteEditorDialog", "Consumido"),
        ("DELETE", "/api/notes/{id}", "Eliminar una nota", "Confirmación", "Consumido"),
        ("POST", "/api/notes/export", "Iniciar exportación asíncrona", "Sin pantalla", "Disponible"),
        ("GET", "/api/admin/users", "Listar usuarios", "Sin pantalla admin", "Disponible"),
        ("PATCH", "/api/admin/users/{id}/role", "Cambiar rol", "Sin pantalla admin", "Disponible"),
    ]
    add_caption(doc, "Tabla 1. Inventario de endpoints y cobertura de interfaz.")
    add_table(doc, ["Método", "Ruta", "Propósito", "Pantalla", "Estado"], endpoint_rows, [850, 2400, 2600, 2000, 1510], 8.2)

    doc.add_heading("1.2 Inventario de pantallas", level=2)
    screen_rows = [
        ("LoginPage", "Acceso", "POST /auth/login", "AppTextField, AppButton", "carga y error"),
        ("RegisterPage", "Alta de cuenta", "POST /auth/register", "AppTextField, AppButton", "carga y error"),
        ("NotesPage", "Listado y CRUD", "GET/POST/PUT/DELETE /notes", "NoteCard, StatePanel, AppButton", "carga, vacío, listo, error"),
        ("NoteEditorDialog", "Crear/editar", "POST/PUT /notes", "AppTextField, AppButton", "validación, guardando, error"),
        ("ComponentCatalogPage", "Documentar el sistema", "Ninguno", "Los cuatro componentes", "variantes y estados"),
    ]
    add_caption(doc, "Tabla 2. Pantallas derivadas de los endpoints y sus estados.")
    add_table(doc, ["Pantalla", "Propósito", "Endpoint", "Componentes", "Estados"], screen_rows, [1500, 1650, 2050, 2600, 1560], 8.4)

    doc.add_heading("2. Sistema de tokens", level=1)
    doc.add_paragraph(
        "Los valores primitivos viven en AppTheme; AppTokens expone nombres semánticos que describen intención. "
        "Los componentes consumen Theme.of(context) y AppTokens.of(context), evitando valores de color y espaciado dispersos."
    )
    color_rows = [
        ("Indigo 700", "primary", "#4338CA", "Acción y foco", "Blanco 7,90:1"),
        ("Blanco", "onPrimary / surface", "#FFFFFF", "Texto inverso / superficie", "Primario 7,90:1"),
        ("Indigo 100", "primarySoft", "#E8E7FF", "Foco y etiqueta suave", "No usado solo como texto"),
        ("Gris 25", "background", "#F6F7FB", "Fondo de aplicación", "Texto 15,17:1"),
        ("Azul tinta", "text", "#182033", "Texto principal", "Fondo 15,17:1"),
        ("Pizarra", "textMuted", "#59657A", "Texto secundario", "Blanco 5,89:1"),
        ("Verde", "success", "#087A55", "Confirmación", "Blanco 5,35:1"),
        ("Ámbar", "warning", "#9A4D00", "Advertencia", "Blanco 6,11:1"),
        ("Rojo", "error", "#B42318", "Error y peligro", "Blanco 6,57:1"),
        ("Gris 300", "outline", "#C8CFDC", "Bordes", "No transmite información solo"),
    ]
    add_caption(doc, "Tabla 3. Tokens de color y contraste WCAG.")
    add_table(doc, ["Primitivo", "Semántico", "Valor", "Uso", "Contraste"], color_rows, [1450, 1800, 1200, 2600, 2310], 8.7)

    token_rows = [
        ("Espaciado", "space1, 2, 3, 4, 6, 8, 12", "4, 8, 12, 16, 24, 32, 48 px"),
        ("Radios", "radiusSm, Md, Lg", "10, 16, 24 px"),
        ("Interacción", "touchTarget", "48 px mínimo"),
        ("Contenido", "contentMaxWidth", "1040 px"),
        ("Tipografía", "display / headline / title / body / label", "36 / 28-22 / 20-16 / 16-14 / 15-13 px"),
    ]
    add_caption(doc, "Tabla 4. Escalas de dimensión y tipografía.")
    add_table(doc, ["Familia", "Token", "Escala"], token_rows, [1700, 3000, 4660], 9)

    doc.add_heading("3. Catálogo de componentes", level=1)
    doc.add_paragraph(
        "Se aplicó la regla de las tres apariciones: los campos y botones aparecen en acceso, registro y editor; "
        "NoteCard representa cada elemento del listado; StatePanel concentra tres estados repetidos de consulta."
    )
    component_rows = [
        ("AppButton", "Acción uniforme y accesible", "label, onPressed, icon, isLoading, variant, expand", "primario, secundario, peligro, deshabilitado, carga"),
        ("AppTextField", "Entrada etiquetada", "controller, label, hint, icon, validator, obscureText, maxLines, enabled", "vacío, valor, error, deshabilitado"),
        ("NoteCard", "Presentar una nota", "note, onEdit, onDelete", "lectura, editable, eliminable, truncado"),
        ("StatePanel", "Explicar estado de consulta", "type, title, message, actionLabel, onAction", "carga, vacío, error"),
    ]
    add_caption(doc, "Tabla 5. Propósito, interfaz pública y estados del catálogo.")
    add_table(doc, ["Componente", "Propósito", "Interfaz pública", "Estados"], component_rows, [1500, 1900, 3600, 2360], 8.6)

    doc.add_heading("4. Pantalla ensamblada y estados", level=1)
    add_labeled_paragraph(doc, "Flujo", "LoginPage / RegisterPage -> ApiService -> NotesPage -> NoteEditorDialog.")
    add_labeled_paragraph(doc, "Delegación", "NoteCard solo emite onEdit y onDelete; NotesPage decide abrir diálogos y llamar a la API.")
    add_labeled_paragraph(doc, "Recuperación", "StatePanel ofrece reintento en error y creación en vacío; RefreshIndicator permite actualización manual.")
    state_rows = [
        ("loading", "Inicio o actualización", "Indicador + texto Cargando tus notas", "Espera o gesto de refrescar"),
        ("empty", "GET devuelve lista vacía", "Icono + explicación + Crear primera nota", "Abrir editor"),
        ("error", "Timeout, red o respuesta inválida", "Icono + mensaje + Reintentar", "Nueva solicitud GET"),
        ("ready", "GET devuelve elementos", "Cuadrícula de NoteCard", "Crear, editar, eliminar, refrescar"),
    ]
    add_caption(doc, "Tabla 6. Máquina de estados de NotesPage.")
    add_table(doc, ["Estado", "Entrada", "Representación", "Salida"], state_rows, [1400, 2400, 3300, 2260], 8.8)

    doc.add_heading("5. Evidencia visual", level=1)
    doc.add_paragraph(
        "Las capturas provienen de app_preview.html, una vista de evidencia construida con los mismos tokens y reglas responsivas del código Flutter. "
        "Los datos son demostrativos; no se presentan como una ejecución del binario Flutter."
    )
    desktop_path = ROOT / "evidence" / "baquero-desktop-1280.png"
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = paragraph.add_run().add_picture(str(desktop_path), width=Inches(6.45))
    set_image_alt(shape, "Baquero escritorio", "Pantalla de notas a 1280 píxeles con métricas y tarjetas en dos columnas.")
    add_caption(doc, "Figura 1. Pantalla ensamblada a 1280 px.")

    add_callout(
        doc,
        "Lectura de escritorio",
        "El ancho de 1280 px activa dos columnas de NoteCard, conserva el contenido en 1040 px y mantiene las acciones principales visibles.",
        accent=PRIMARY,
        fill=PRIMARY_SOFT,
    )
    doc.add_page_break()
    doc.add_heading("5.1 Comparación móvil y fuente ampliada", level=2)

    mobile_path = ROOT / "evidence" / "baquero-mobile-390.png"
    large_path = ROOT / "evidence" / "baquero-mobile-390-font-150.png"
    comparison = doc.add_table(rows=2, cols=2)
    comparison.style = "Table Grid"
    set_table_geometry(comparison, [4680, 4680])
    repeat_header(comparison.rows[0])
    for index, label in enumerate(("390 px - fuente normal", "390 px - fuente al 150 %")):
        cell = comparison.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_GRAY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(label)
        set_font(run, size=9, bold=True, color=INK)
    paths = (mobile_path, large_path)
    alts = (
        ("Baquero móvil", "Pantalla de notas a 390 píxeles con una columna y controles accesibles."),
        ("Baquero móvil con fuente ampliada", "Pantalla a 390 píxeles con fuente del sistema al 150 por ciento y sin desplazamiento horizontal."),
    )
    for index, source in enumerate(paths):
        cell = comparison.rows[1].cells[index]
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shape = paragraph.add_run().add_picture(str(source), width=Inches(1.95))
        set_image_alt(shape, alts[index][0], alts[index][1])
    for row in comparison.rows:
        prevent_row_split(row)
    add_caption(doc, "Figuras 2 y 3. Pantalla móvil normal y verificación con fuente del sistema ampliada.")
    doc.add_page_break()

    doc.add_heading("6. Verificación de accesibilidad", level=1)
    accessibility_rows = [
        ("Contraste de texto", "AA >= 4,5:1", "5,35:1 a 15,17:1", "Cumple"),
        ("Objetivos táctiles", "48 x 48 px", "touchTarget e IconButton", "Cumple"),
        ("Etiquetas", "Nombre accesible", "Semantics, labelText y Tooltip", "Cumple por inspección"),
        ("Color", "No ser único medio", "Icono + título + mensaje + acción", "Cumple"),
        ("Anchos", "1280 y 390 px", "Capturas sin solapamiento", "Cumple"),
        ("Fuente ampliada", "150 %", "Desbordamiento detectado y corregido con corte seguro", "Cumple"),
        ("Código Flutter", "Prueba de widget", "Pruebas preparadas; SDK no disponible", "Pendiente de ejecución"),
    ]
    add_caption(doc, "Tabla 7. Matriz de verificación y resultados.")
    add_table(doc, ["Criterio", "Umbral", "Evidencia", "Resultado"], accessibility_rows, [1900, 1800, 3800, 1860], 8.8)
    add_callout(
        doc,
        "Corrección aplicada",
        "La prueba a 150 % descubrió desplazamiento horizontal causado por un correo largo. Se añadió overflow-wrap:anywhere al encabezado y se repitió la captura: scrollWidth quedó igual al ancho útil (375 px).",
        accent=WARNING,
        fill="FFF4E8",
    )

    doc.add_heading("7. Verificación técnica y limitaciones", level=1)
    verification_rows = [
        ("Compilación Python", "python -m py_compile backend/app.py", "Correcta"),
        ("Integración backend", "registro / crear / listar / actualizar / eliminar", "BACKEND_SMOKE_OK"),
        ("Compatibilidad JWT", "sub como String y conversión a int para SQLite", "Fallo detectado y corregido"),
        ("Pruebas Flutter", "test/api_service_test.dart y test/widget_test.dart", "Preparadas, no ejecutadas"),
        ("Capturas", "1280 px, 390 px y 390 px al 150 %", "Revisadas visualmente"),
    ]
    add_caption(doc, "Tabla 8. Evidencia técnica ejecutada.")
    add_table(doc, ["Prueba", "Alcance", "Resultado"], verification_rows, [2000, 4800, 2560], 9)
    add_labeled_paragraph(doc, "Limitación del entorno", "No había Flutter ni Dart en las rutas disponibles; por eso no se afirma que flutter test o flutter analyze hayan sido ejecutados.")
    add_labeled_paragraph(doc, "Repositorio local", str(ROOT))
    add_labeled_paragraph(doc, "Repositorio remoto", "[PENDIENTE: pegar URL GitHub/GitLab con permisos de lectura]. La copia entregada no contiene metadatos .git válidos.")

    doc.add_heading("8. Registro de uso de inteligencia artificial", level=1)
    ai_rows = [
        ("Herramienta", "OpenAI Codex"),
        ("Solicitud", "Implementar el taller de componentes reutilizables y entregar el informe en Word."),
        ("Aportes", "Inventario, arquitectura, tokens, componentes, integración HTTP, pruebas, capturas e informe."),
        ("Cambios humanos pendientes", "Completar portada, proporcionar URL del repositorio y ejecutar Flutter en una estación con SDK."),
        ("Verificación", "Prueba backend aislada, cálculo WCAG, inspección responsiva y render visual del DOCX."),
    ]
    add_caption(doc, "Tabla 9. Trazabilidad del apoyo de IA.")
    add_table(doc, ["Campo", "Registro"], ai_rows, [2100, 7260], 9.2)

    doc.add_heading("9. Conclusiones", level=1)
    doc.add_paragraph(
        "La implementación cumple el objetivo central del taller: transforma decisiones visuales en tokens, abstrae componentes por uso repetido, "
        "mantiene la API fuera de los componentes, ensambla una pantalla real y hace visibles los estados operativos. La prueba de integración también "
        "permitió corregir una incompatibilidad JWT que habría bloqueado el cliente. Antes de la entrega final solo faltan los datos académicos de portada, "
        "la URL remota y la ejecución de la suite Flutter en un equipo con SDK."
    )

    doc.add_heading("Anexo A. Código fuente seleccionable", level=1)
    doc.add_paragraph(
        "Los siguientes listados son texto seleccionable, no capturas. Se incluyen completos y con numeración editorial de línea para facilitar la revisión."
    )

    components = [
        (
            "A.1 AppButton",
            ROOT / "lib" / "components" / "app_button.dart",
            [
                ("Contrato", "Recibe contenido, variante, carga y devolución de llamada; no decide navegación ni persistencia."),
                ("Tema", "Obtiene colores y dimensiones desde AppTokens, por lo que cambia junto con el tema."),
                ("Accesibilidad", "Expone Semantics, estado habilitado y un objetivo táctil mínimo de 48 px."),
            ],
        ),
        (
            "A.2 AppTextField",
            ROOT / "lib" / "components" / "app_text_field.dart",
            [
                ("Contrato", "Parametriza etiqueta, controlador, validación, teclado, líneas y estado habilitado."),
                ("Composición", "Delega la regla de validación al formulario padre y conserva una presentación uniforme."),
                ("Accesibilidad", "Combina Semantics con labelText visible; no depende de placeholder como única etiqueta."),
            ],
        ),
        (
            "A.3 NoteCard",
            ROOT / "lib" / "components" / "note_card.dart",
            [
                ("Modelo", "Recibe una Note inmutable y solo presenta sus datos."),
                ("Delegación", "onEdit y onDelete permiten reutilizar la tarjeta sin conocer rutas ni backend."),
                ("Robustez", "Trunca contenido largo, adapta autor opcional y mantiene controles de 48 px con Tooltip."),
            ],
        ),
        (
            "A.4 StatePanel",
            ROOT / "lib" / "components" / "state_panel.dart",
            [
                ("Estados", "Un enum concentra carga, vacío y error sin duplicar estructura en las pantallas."),
                ("Recuperación", "actionLabel y onAction permiten crear o reintentar sin acoplar la acción."),
                ("Accesibilidad", "La región viva anuncia título y mensaje, y cada estado combina icono, texto y color."),
            ],
        ),
    ]

    for index, (heading, source_path, explanations) in enumerate(components):
        doc.add_heading(heading, level=2)
        for label, text in explanations:
            add_labeled_paragraph(doc, label, text)
        add_code_block(doc, source_path)

    doc.core_properties.title = "Sistema de componentes reutilizables para Baquero Notes"
    doc.core_properties.subject = "Informe técnico de Flutter y API Flask"
    doc.core_properties.author = "[COMPLETAR NOMBRE]"
    doc.core_properties.keywords = "Flutter, componentes, tokens, accesibilidad, Flask, API"
    doc.core_properties.comments = "Generado con apoyo de OpenAI Codex; requiere revisión académica final."

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
